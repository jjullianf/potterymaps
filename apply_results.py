"""
Wendet die in results.json gesammelten Recherche-Ergebnisse final auf
keramik-bemalen-schweiz-schema.xlsx an:
- Fuellt die kept-Zeilen mit den recherchierten Feldern
- Fuegt Spalte 'hat_instagram' hinzu (Spalte R)
- Verschiebt excluded-Zeilen in einen neuen Tab 'Nachtraeglich ausgeschlossen'
  und entfernt sie aus 'Studios CH'
"""
import json
import openpyxl
from openpyxl.styles import Font

RESULTS_PATH = "/private/tmp/claude-501/-Users-julianfrick-Desktop-keramik-directory/66778423-d814-45c6-8ec0-1bc3dbb2e937/scratchpad/results.json"
XLSX_PATH = "keramik-bemalen-schweiz-schema.xlsx"

COLS = {
    "opening_hours": 7, "tags": 8, "price_mug": 9, "price_plate": 10,
    "studio_fee": 11, "price_note": 12, "pickup_time": 13, "special_events": 14,
    "description": 15, "hat_instagram": 18,
}


def main():
    data = json.load(open(RESULTS_PATH, encoding="utf-8"))
    kept = data["kept"]
    excluded = data["excluded"]

    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb["Studios CH"]

    if ws.cell(row=1, column=18).value != "hat_instagram":
        ws.cell(row=1, column=18, value="hat_instagram")
        ws.cell(row=1, column=18).font = Font(bold=True)

    for row_str, fields in kept.items():
        row = int(row_str)
        for field, value in fields.items():
            col = COLS[field]
            ws.cell(row=row, column=col, value=value)

    if "Nachtraeglich ausgeschlossen" in wb.sheetnames:
        del wb["Nachtraeglich ausgeschlossen"]
    ex_ws = wb.create_sheet("Nachtraeglich ausgeschlossen")
    headers = ["name", "city", "address", "website", "begruendung"]
    for c, h in enumerate(headers, start=1):
        ex_ws.cell(row=1, column=c, value=h)
        ex_ws.cell(row=1, column=c).font = Font(bold=True)

    ex_row = 2
    for row_str, reason in excluded.items():
        row = int(row_str)
        name = ws.cell(row=row, column=2).value
        address = ws.cell(row=row, column=3).value
        city = ws.cell(row=row, column=4).value
        website = ws.cell(row=row, column=6).value
        ex_ws.cell(row=ex_row, column=1, value=name)
        ex_ws.cell(row=ex_row, column=2, value=city)
        ex_ws.cell(row=ex_row, column=3, value=address)
        ex_ws.cell(row=ex_row, column=4, value=website)
        ex_ws.cell(row=ex_row, column=5, value=reason)
        ex_row += 1

    for col, width in zip("ABCDE", [35, 20, 35, 35, 90]):
        ex_ws.column_dimensions[col].width = width

    rows_to_delete = sorted((int(r) for r in excluded.keys()), reverse=True)
    for row in rows_to_delete:
        ws.delete_rows(row)

    wb.save(XLSX_PATH)
    print(f"Fertig: {len(kept)} Zeilen aktualisiert, {len(excluded)} Zeilen ausgeschlossen und verschoben.")
    print(f"Studios CH hat jetzt {ws.max_row - 1} Zeilen (ohne Header).")


if __name__ == "__main__":
    main()
