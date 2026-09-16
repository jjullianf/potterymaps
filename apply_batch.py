"""
Traegt die von Claude recherchierten Werte fuer einen Batch von Zeilen in
keramik-bemalen-schweiz-schema.xlsx ein. BATCH_DATA und EXCLUDED werden vor
jedem Lauf angepasst. Ausgeschlossene Zeilen werden aus 'Studios CH' entfernt
und in den Tab 'Nachtraeglich ausgeschlossen' verschoben.
"""
import openpyxl
from openpyxl.styles import Font
from openpyxl.comments import Comment

XLSX_PATH = "keramik-bemalen-schweiz-schema.xlsx"

COLS = {
    "opening_hours": 7, "tags": 8, "price_mug": 9, "price_plate": 10,
    "studio_fee": 11, "price_note": 12, "pickup_time": 13, "special_events": 14,
    "description": 15, "hat_instagram": 18,
}

# row -> {field: value}
BATCH_DATA = {}

# row -> reason string
EXCLUDED = {}


def main():
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb["Studios CH"]

    if ws.cell(row=1, column=18).value != "hat_instagram":
        ws.cell(row=1, column=18, value="hat_instagram")
        ws.cell(row=1, column=18).font = Font(bold=True)

    if "Nachtraeglich ausgeschlossen" not in wb.sheetnames:
        ex_ws = wb.create_sheet("Nachtraeglich ausgeschlossen")
        ex_ws.cell(row=1, column=1, value="name")
        ex_ws.cell(row=1, column=2, value="city")
        ex_ws.cell(row=1, column=3, value="website")
        ex_ws.cell(row=1, column=4, value="begruendung")
        for c in range(1, 5):
            ex_ws.cell(row=1, column=c).font = Font(bold=True)
    else:
        ex_ws = wb["Nachtraeglich ausgeschlossen"]

    for row, fields in BATCH_DATA.items():
        for field, value in fields.items():
            col = COLS[field]
            ws.cell(row=row, column=col, value=value)

    if EXCLUDED:
        next_ex_row = ex_ws.max_row + 1
        rows_to_delete = sorted(EXCLUDED.keys(), reverse=True)
        for row in sorted(EXCLUDED.keys()):
            name = ws.cell(row=row, column=2).value
            city = ws.cell(row=row, column=4).value
            website = ws.cell(row=row, column=6).value
            ex_ws.cell(row=next_ex_row, column=1, value=name)
            ex_ws.cell(row=next_ex_row, column=2, value=city)
            ex_ws.cell(row=next_ex_row, column=3, value=website)
            ex_ws.cell(row=next_ex_row, column=4, value=EXCLUDED[row])
            next_ex_row += 1
        for row in rows_to_delete:
            ws.delete_rows(row)

    wb.save(XLSX_PATH)
    print(f"Batch angewendet: {len(BATCH_DATA)} Zeilen aktualisiert, {len(EXCLUDED)} ausgeschlossen.")


if __name__ == "__main__":
    main()
