"""
Crawlt fuer jedes der 54 Keramik-Bemal-Studios die Website (Startseite + relevante
Unterseiten wie Preise/Kurse/Kontakt/FAQ) und speichert den gesammelten Text plus
gefundene Social-Media-Links als JSON-Datei pro Zeile ab. Dient als Recherche-Basis,
die anschliessend manuell/durch Claude ausgewertet wird - keine automatische
Klassifikation oder Texterstellung hier.
"""

import asyncio
import json
import re
import sys
from pathlib import Path

from crawl4ai import AsyncWebCrawler, BrowserConfig, CrawlerRunConfig, CacheMode
import openpyxl

XLSX_PATH = Path("keramik-bemalen-schweiz-schema.xlsx")
OUT_DIR = Path("/private/tmp/claude-501/-Users-julianfrick-Desktop-keramik-directory/66778423-d814-45c6-8ec0-1bc3dbb2e937/scratchpad/studio-dumps")
OUT_DIR.mkdir(parents=True, exist_ok=True)

MAX_SUBPAGES = 6
PAGE_TIMEOUT_MS = 25_000

SUBPAGE_HINTS = [
    "preis", "price", "pricing", "tarif", "kurse", "kurs", "angebot", "workshop",
    "offen", "walk", "events", "event", "faq", "kontakt", "contact", "about",
    "ueber", "über", "info", "buchen", "book", "termine", "oeffnungszeiten",
    "öffnungszeiten", "hours", "cours", "prix", "atelier", "prezzi", "corso",
]

BROWSER_CONFIG = BrowserConfig(headless=True, verbose=False)
RUN_CONFIG = CrawlerRunConfig(
    cache_mode=CacheMode.BYPASS,
    page_timeout=PAGE_TIMEOUT_MS,
    word_count_threshold=0,
    semaphore_count=5,
    delay_before_return_html=2.5,
    verbose=False,
)


def clean_url(url: str) -> str:
    if "%3F" in url:
        url = url.split("%3F")[0]
    return url


def pick_subpages(links: dict, base_href: str) -> list[str]:
    internal = links.get("internal", []) if isinstance(links, dict) else []
    seen = {base_href}
    picked = []
    for link in internal:
        href = (link.get("href") or "").split("#")[0]
        text = (link.get("text") or "").lower()
        if not href or href in seen:
            continue
        combined = f"{href.lower()} {text}"
        if any(hint in combined for hint in SUBPAGE_HINTS):
            seen.add(href)
            picked.append(href)
        if len(picked) >= MAX_SUBPAGES:
            break
    return picked


def find_social_links(links: dict) -> dict:
    external = links.get("external", []) if isinstance(links, dict) else []
    socials = {}
    for link in external:
        href = (link.get("href") or "")
        low = href.lower()
        if "instagram.com" in low and "instagram" not in socials:
            socials["instagram"] = href
        elif "facebook.com" in low and "facebook" not in socials:
            socials["facebook"] = href
    return socials


async def fetch_one(crawler: AsyncWebCrawler, row_idx: int, name: str, website: str) -> dict:
    website = clean_url(website)
    data = {"row": row_idx, "name": name, "website": website, "pages": {}, "social": {}, "error": None}

    if not website:
        data["error"] = "keine Website angegeben"
        return data

    try:
        home = await crawler.arun(url=website, config=RUN_CONFIG)
    except Exception as exc:  # noqa: BLE001
        data["error"] = f"Fehler beim Laden: {exc}".splitlines()[0][:300]
        return data

    if not home.success:
        data["error"] = f"nicht erreichbar: {(home.error_message or '').splitlines()[0][:300]}"
        return data

    data["pages"][website] = home.markdown or ""
    data["social"] = find_social_links(home.links or {})

    subpages = pick_subpages(home.links or {}, website)
    for sub_url in subpages:
        try:
            sub = await crawler.arun(url=sub_url, config=RUN_CONFIG)
        except Exception:
            continue
        if sub.success and sub.markdown:
            data["pages"][sub_url] = sub.markdown
            sub_social = find_social_links(sub.links or {})
            for k, v in sub_social.items():
                data["social"].setdefault(k, v)

    return data


async def main():
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb["Studios CH"]

    rows = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=2).value
        website = (ws.cell(row=r, column=6).value or "").strip()
        if name is None:
            continue
        rows.append((r, name, website))

    print(f"{len(rows)} Studios zu crawlen")

    async with AsyncWebCrawler(config=BROWSER_CONFIG) as crawler:
        batch_size = 10
        for i in range(0, len(rows), batch_size):
            batch = rows[i : i + batch_size]
            print(f"\n=== Crawl-Batch {i // batch_size + 1} ({len(batch)} Studios) ===")
            tasks = [fetch_one(crawler, r, name, website) for r, name, website in batch]
            results = await asyncio.gather(*tasks)
            for res in results:
                out_path = OUT_DIR / f"row_{res['row']:03d}.json"
                out_path.write_text(json.dumps(res, ensure_ascii=False, indent=2), encoding="utf-8")
                status = res["error"] or f"{len(res['pages'])} Seiten, social={list(res['social'].keys())}"
                print(f"  row {res['row']:3d} | {res['name'][:50]:50s} | {status}")

    print("\nFertig. Dumps liegen in", OUT_DIR)


if __name__ == "__main__":
    asyncio.run(main())
